"""Run AFTER train_model.py:  python build_report.py   -> outputs/Project_Report.docx (numbers come from metrics.json)"""
import json, pandas as pd
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from credit_utils import *

m = json.load(open("models/metrics.json")); c = m["cleaning"]
df = pd.read_csv("data/clean_data.csv")
gr = df.groupby("loan_grade").loan_status.mean(); ho = df.groupby("person_home_ownership").loan_status.mean()
best = m["best_model"]; cmp_ = m["comparison"]; bands = m["risk_bands"]

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
def H(t, l=1): doc.add_heading(t, l)
def P(t): doc.add_paragraph(t)
def B(items):
    for i in items: doc.add_paragraph(i, style="List Bullet")
def IMG(f, cap, w=6.2):
    doc.add_picture("outputs/charts/" + f, width=Inches(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(cap); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic = True
def TABLE(rows, header):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"
    for i, h in enumerate(header): t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r): cells[i].text = str(v)

t = doc.add_heading("Credit Risk Analysis System", 0)
P("Data Science Capstone Project - SkillOrbit")
P("Student: Ananya Kumar Srivastava | MNIT Jaipur, B.Tech Civil Engineering (2029)")

H("1. Introduction")
P("Banks and NBFCs must decide whether an applicant will repay a loan. A wrong approval causes a loss; a wrong rejection "
  "loses a good customer. This project builds a data science application that estimates each applicant's probability of "
  "default (PD) from financial and credit-history data, groups applicants into Low / Medium / High risk, shows the results "
  "on an interactive dashboard, and recommends one of three actions: Approve Loan, Review Manually or Reject Application.")

H("2. Objectives")
B(["Analyze customer financial data and find what drives default.",
   "Identify low-risk and high-risk applicants.",
   "Predict loan approval probability with a machine learning classifier.",
   "Generate risk insights through a dashboard and visualize risk patterns.",
   "Produce an automatic, explainable loan recommendation."])

H("3. Dataset")
P(f"The project uses the Kaggle 'Credit Risk Dataset' schema with 12 columns: age, income, home ownership, employment length, "
  f"loan purpose, loan grade, loan amount, interest rate, loan-to-income ratio, prior default flag and credit-history length. "
  f"The target column is loan_status (1 = defaulted, 0 = repaid). The raw file had {c['rows_raw']:,} rows.")

H("4. Methodology")
H("4.1 Data cleaning and preprocessing (Module 1)", 2)
B([f"Duplicates removed: {c['duplicates_removed']:,}.",
   f"Impossible ages (e.g. 144) removed: {c['impossible_age_removed']}; impossible employment lengths (e.g. 123 years) removed: {c['impossible_emp_length_removed']}.",
   f"Missing values: {c['missing_emp_length']:,} in employment length and {c['missing_loan_int_rate']:,} in interest rate - filled with the median "
   "inside the model pipeline (fitted on training data only, so no data leakage).",
   "Categorical columns encoded with one-hot encoding; numeric columns normalised with StandardScaler.",
   f"Final clean dataset: {c['rows_clean']:,} rows. Split 80/20 into train ({m['n_train']:,}) and test ({m['n_test']:,}), stratified on the target."])
H("4.2 Exploratory data analysis (Module 2)", 2)
B([f"Only about {df.loan_status.mean():.0%} of loans default - the classes are imbalanced, so accuracy alone is misleading.",
   f"Default rate rises steadily with loan grade: grade A {gr['A']:.0%} -> grade D {gr['D']:.0%} -> grade G {gr['G']:.0%}.",
   f"Renters default more ({ho['RENT']:.0%}) than home owners ({ho['OWN']:.0%}).",
   "Loan-to-income ratio and interest rate are clearly higher for defaulters; income is lower."])
IMG("01_target_and_income.png", "Figure 1: Class balance and income distribution")
IMG("02_default_rates.png", "Figure 2: Default rate by grade, home ownership and loan purpose")
IMG("03_boxplots.png", "Figure 3: Key numeric features versus default")
IMG("04_correlation.png", "Figure 4: Correlation heatmap", 4.8)
H("4.3 Model building (Module 3)", 2)
P("Three classifiers recommended by the brief were trained inside one scikit-learn Pipeline (preprocessing + model), using "
  "class weighting to handle the imbalance: Logistic Regression, Decision Tree (depth 6) and Random Forest (300 trees). "
  "They were compared on the untouched test set.")
TABLE([[n] + [f"{v:.3f}" for v in r.values()] for n, r in cmp_.items()], ["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC"])
P("")
P(f"The best model by ROC-AUC was {best}. Because class weighting inflates predicted probabilities, the selected model was "
  f"wrapped in isotonic probability calibration so that a prediction of 20% really means about 20% of such applicants default "
  f"(calibrated test ROC-AUC = {m['calibrated_roc_auc']:.3f}).")
IMG("05_roc_curves.png", "Figure 5: ROC curves of the three models", 4.6)
IMG("06_model_diagnostics.png", "Figure 6: Confusion matrix and the main drivers of risk")

H("5. Risk score and recommendation system (Modules 4-5)")
P(f"The calibrated PD is converted into a risk band and a decision: PD below {APPROVE_BELOW:.0%} = Low Risk -> Approve Loan; "
  f"{APPROVE_BELOW:.0%} to {REJECT_FROM:.0%} = Medium Risk -> Review Manually; {REJECT_FROM:.0%} and above = High Risk -> Reject Application. "
  "Thresholds are constants that a bank can tune to its risk appetite.")
TABLE([[k, f"{v['applicants']:,}", f"{v['actual_default_rate']:.1%}", f"{v['avg_pd']:.1%}"] for k, v in bands.items()],
      ["Risk band", "Test applicants", "Actual default rate", "Average predicted PD"])
P("")
P("The bands are well separated and the predicted PD closely matches the actual default rate, which validates the score.")
IMG("07_risk_bands.png", "Figure 7: Predicted PD distribution and actual default rate per band")

H("6. Dashboard")
P("The Streamlit application has three tabs: (1) Risk Predictor - enter an applicant, get a risk gauge, default and approval "
  "probability, risk category, recommendation and plain-language risk flags; (2) Portfolio Dashboard - KPIs, risk distribution, "
  "recommendation counts, default rate by grade and age, loan trends, with filters; (3) Model Performance - metrics and diagnostics.")

H("7. Deployment")
P("The code is hosted in a GitHub repository and deployed on Streamlit Community Cloud. Dependencies are listed in requirements.txt. "
  "Live link: <paste your Streamlit URL here>. GitHub: <paste your repo URL here>.")

H("8. Limitations and future work")
B(["Public dataset only - no real banking integration, as per scope.",
   "Thresholds are illustrative; real lenders tune them using cost of default versus cost of lost business.",
   "Model uses no fairness audit; a production system should test for bias across demographic groups.",
   "Future work: SHAP explanations, hyper-parameter tuning, gradient boosting, MySQL storage, batch CSV scoring."])

H("9. Conclusion")
P("The project delivers an end-to-end data science workflow - cleaning, EDA, classification, calibrated risk scoring, a dashboard "
  "and an automated recommendation - that behaves like a simplified credit-risk platform used by banks and fintech companies.")
doc.save("outputs/Project_Report.docx"); print("saved outputs/Project_Report.docx")
